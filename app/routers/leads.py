"""Отдел продаж → «Прозвон».

Импорт спарсенного Excel/CSV со списком точек (кофейни, кондитерские и т.п.),
автоматическое выявление сетевых / одиночных точек, извлечение соцсетей и
удобная таблица для прозвона с отметками статуса.
"""
import csv
import io
import json
import os
import re
import tempfile
import threading
import time
import uuid
from datetime import datetime, date

from fastapi import APIRouter, Request, Depends, Form, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func

from app.database import get_db
from app.auth import login_required, role_required
from app.models import SalesLead, LeadCall, User, Counterparty
from app.utils.leads_utils import (
    _normalize_brand, _extract_socials, _ensure_scheme, _norm_phone,
    PHONE_PATTERN, URL_PATTERN,
)

# Поля, которые можно сопоставлять колонкам (для превью-маппинга)
MAPPABLE_FIELDS = {
    "name": "Название", "phone": "Телефон", "email": "Email",
    "city": "Город", "district": "Район", "address": "Адрес", "category": "Рубрика",
    "contact_person": "Контактное лицо", "website": "Сайт",
    "vk": "VK", "instagram": "Instagram", "telegram": "Telegram", "whatsapp": "WhatsApp",
}
_TMP_DIR = os.path.join(tempfile.gettempdir(), "tms_leads_import")

router = APIRouter(prefix="/leads", tags=["leads"])
templates = Jinja2Templates(directory="app/templates")

LEAD_STATUSES = {
    "new":        "Новый",
    "callback":   "Перезвонить",
    "no_answer":  "Не дозвонился",
    "interested": "Заинтересован",
    "thinking":   "Думает",
    "refused":    "Отказ",
    "deal":       "Продажа / договор",
    "invalid":    "Невалид",
}
STATUS_COLORS = {
    "new": "secondary", "callback": "info", "no_answer": "warning",
    "interested": "primary", "thinking": "info", "refused": "danger",
    "deal": "success", "invalid": "dark",
}

# ── Эвристики сопоставления колонок ──────────────────────────────────────────
COLUMN_HINTS = {
    "name":    ["назван", "наименован", "компан", "организац", "заведен", "точка",
                "name", "title", "фирма", "объект", "бренд"],
    "phone":   ["телефон", "тел.", "тел ", "phone", "моб", "контактный тел", "номер"],
    "email":   ["e-mail", "email", "почта", "mail"],
    "city":     ["город", "city", "населен"],
    "district": ["район", "district", "р-н", "округ"],
    "address":  ["адрес", "address", "местоположен", "улиц"],
    "category":["рубрик", "категор", "вид деятельн", "сфера", "тип заведен",
                "профиль", "catalog", "отрасл"],
    "contact_person": ["контактное лицо", "контактн", "контакт", "фио", "директор", "руковод"],
    "website": ["сайт", "site", "web", "url", "домен", "веб"],
    "vk":        ["вконтакте", "vk", "вк"],
    "instagram": ["instagram", "инстаграм", "инст"],
    "telegram":  ["telegram", "телеграм", "tg"],
    "whatsapp":  ["whatsapp", "ватсап", "вотсап", "wa "],
}

# Соцсети / онлайн — извлечение по содержимому ячеек
SOCIAL_PATTERNS = {
    "vk":        re.compile(r"(?:https?://)?(?:www\.)?(?:vk\.com|vk\.ru|m\.vk\.com)/[\w.\-/]+", re.I),
    "instagram": re.compile(r"(?:https?://)?(?:www\.)?instagram\.com/[\w.\-/]+", re.I),
    "telegram":  re.compile(r"(?:https?://)?(?:www\.)?(?:t\.me|telegram\.me)/[\w.\-/]+", re.I),
    "whatsapp":  re.compile(r"(?:https?://)?(?:www\.)?(?:wa\.me|api\.whatsapp\.com)/[\w.\-/?=&]+", re.I),
}
# URL_PATTERN, PHONE_PATTERN, _normalize_brand, _extract_socials, _ensure_scheme, _norm_phone — из leads_utils

# Регулярка для авто-извлечения района из строки адреса:
# «Центральный р-н», «р-н Советский», «Октябрьский район», «район Коминтерна»
_DISTRICT_RE = re.compile(
    r'([А-ЯЁа-яё][а-яёА-ЯЁ\w\s\-]{1,40}?)\s+(?:р[-‐–\s]?н\.?|район)'
    r'|(?:р[-‐–\s]?н\.?|район)\s+([А-ЯЁа-яё][а-яёА-ЯЁ\w\s\-]{1,40})',
    re.UNICODE,
)


def _extract_district(address: str) -> str:
    """Пытается извлечь название района из строки адреса. Возвращает '' если не найдено."""
    if not address:
        return ""
    m = _DISTRICT_RE.search(address)
    if not m:
        return ""
    raw = (m.group(1) or m.group(2) or "").strip().rstrip(",;. ")
    return raw[:150] if raw else ""


def _match_columns(headers: list[str]) -> dict:
    """Сопоставляет индексы колонок с полями по подсказкам."""
    mapping = {}
    used = set()
    norm = [(i, (h or "").strip().lower()) for i, h in enumerate(headers)]
    for field, hints in COLUMN_HINTS.items():
        for i, h in norm:
            if i in used or not h:
                continue
            if any(hint in h for hint in hints):
                mapping[field] = i
                used.add(i)
                break
    return mapping


# _extract_socials — из leads_utils


def _clean(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


# _ensure_scheme — из leads_utils


def _parse_rows(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    """Возвращает (headers, rows) из xlsx или csv."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        text = None
        for enc in ("utf-8-sig", "cp1251", "utf-8", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        text = text or data.decode("utf-8", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delim)
        rows = [[_clean(c) for c in r] for r in reader]

    rows = [r for r in rows if any(c for c in r)]   # выбросить пустые строки
    if not rows:
        return [], []
    headers = rows[0]
    return headers, rows[1:]


@router.get("/", response_class=HTMLResponse)
@login_required
async def list_leads(
    request: Request,
    q: str = "", status: str = "", kind: str = "",
    category: str = "", district: str = "", assigned: str = "", source: str = "", due: str = "",
    db: Session = Depends(get_db),
):
    query = db.query(SalesLead).filter(SalesLead.is_active == True)
    if q:
        like = f"%{q}%"
        query = query.filter(
            SalesLead.name.ilike(like) | SalesLead.phone.ilike(like) |
            SalesLead.address.ilike(like) | SalesLead.category.ilike(like)
        )
    if status:
        query = query.filter(SalesLead.call_status == status)
    if kind == "network":
        query = query.filter(SalesLead.is_network == True)
    elif kind == "single":
        query = query.filter(SalesLead.is_network == False)
    if category:
        query = query.filter(SalesLead.category.ilike(f"%{category}%"))
    if district:
        query = query.filter(SalesLead.district.ilike(f"%{district}%"))
    if assigned == "none":
        query = query.filter(SalesLead.assigned_to_id.is_(None))
    elif assigned:
        query = query.filter(SalesLead.assigned_to_id == int(assigned))
    if source:
        query = query.filter(SalesLead.source_file == source)
    if due == "today":
        query = query.filter(SalesLead.callback_at <= date.today(),
                             SalesLead.callback_at.isnot(None))

    leads = query.order_by(
        SalesLead.is_network.desc(), SalesLead.brand, SalesLead.name
    ).all()

    # Сводка по всем активным (без учёта фильтров — для карточек)
    base = db.query(SalesLead).filter(SalesLead.is_active == True)
    stats = {
        "total":      base.count(),
        "networks":   base.filter(SalesLead.is_network == True).count(),
        "called":     base.filter(SalesLead.call_status != "new").count(),
        "interested": base.filter(SalesLead.call_status.in_(["interested", "thinking"])).count(),
        "deals":      base.filter(SalesLead.call_status == "deal").count(),
        "due_today":  base.filter(SalesLead.callback_at <= date.today(),
                                  SalesLead.callback_at.isnot(None)).count(),
    }
    users = db.query(User).filter(
        User.is_active == True, User.role == "sales"
    ).order_by(User.full_name).all()
    sources = [s[0] for s in db.query(SalesLead.source_file)
               .filter(SalesLead.is_active == True, SalesLead.source_file.isnot(None))
               .distinct().all()]
    categories = sorted({c[0] for c in db.query(SalesLead.category)
                         .filter(SalesLead.is_active == True, SalesLead.category.isnot(None)).all() if c[0]})
    districts = sorted({d[0] for d in db.query(SalesLead.district)
                        .filter(SalesLead.is_active == True, SalesLead.district.isnot(None)).all() if d[0]})

    # ── Группировка: сетевые точки сворачиваем в одну группу по бренду ──
    from collections import OrderedDict
    net_groups: "OrderedDict[str, list]" = OrderedDict()
    groups = []
    for l in leads:
        if l.is_network and l.brand:
            net_groups.setdefault(l.brand, []).append(l)
        else:
            groups.append({"kind": "single", "lead": l})
    # вставляем сети в начало (leads уже отсортированы network-first)
    net_items = []
    for brand, items in net_groups.items():
        net_items.append({
            "kind": "network", "brand": brand,
            "name": items[0].name, "count": len(items), "leads": items,
            "category": items[0].category,
            "cities": sorted({i.city for i in items if i.city}),
            "done": sum(1 for i in items if i.call_status != "new"),
            "deals": sum(1 for i in items if i.call_status == "deal"),
        })
    groups = net_items + groups

    return templates.TemplateResponse(request, "leads/list.html", {
        "leads": leads, "groups": groups, "stats": stats, "users": users,
        "sources": sources, "categories": categories, "districts": districts,
        "statuses": LEAD_STATUSES, "status_colors": STATUS_COLORS,
        "q": q, "status": status, "kind": kind, "category": category,
        "district": district, "assigned": assigned, "source": source, "due": due,
    })


# _norm_phone — из leads_utils


def _build_lead(cells, cols, headers) -> dict | None:
    def g(field):
        idx = cols.get(field)
        return _clean(cells[idx]) if idx is not None and idx < len(cells) else ""

    name_idx = cols.get("name", 0)
    name = _clean(g("name") or (cells[name_idx] if name_idx < len(cells) else ""))
    if not name:
        return None
    socials = _extract_socials(cells)
    phone = g("phone")
    if not phone:
        m = PHONE_PATTERN.search(" ".join(cells))
        phone = m.group(0).strip() if m else ""
    # бренд для группировки сетей; если нормализация всё «съела» — берём само название,
    # чтобы идентичные названия гарантированно попадали в одну сеть
    brand = _normalize_brand(name) or name.lower().strip()
    address_val = g("address")
    district_val = g("district") or _extract_district(address_val)
    return {
        "name": name[:300],
        "brand": brand[:300],
        "category": g("category")[:150],
        "city": g("city")[:150],
        "district": district_val[:150],
        "address": address_val[:500],
        "phone": phone[:150],
        "email": g("email")[:150],
        "contact_person": g("contact_person")[:150],
        "website":   _ensure_scheme(g("website") or socials.get("website", ""))[:500],
        "vk":        _ensure_scheme(g("vk") or socials.get("vk", ""))[:500],
        "instagram": _ensure_scheme(g("instagram") or socials.get("instagram", ""))[:500],
        "telegram":  _ensure_scheme(g("telegram") or socials.get("telegram", ""))[:500],
        "whatsapp":  _ensure_scheme(g("whatsapp") or socials.get("whatsapp", ""))[:500],
        "raw": json.dumps(dict(zip(headers, cells)), ensure_ascii=False),
    }


@router.get("/import", response_class=HTMLResponse)
@login_required
async def import_form(request: Request):
    return templates.TemplateResponse(request, "leads/import.html",
                                      {"step": "upload", "result": None})


@router.post("/import", response_class=HTMLResponse)
@login_required
async def import_preview(request: Request, file: UploadFile = File(...)):
    """Шаг 1 — парсим, сохраняем во временный файл, показываем превью маппинга."""
    data = await file.read()
    try:
        headers, rows = _parse_rows(file.filename, data)
    except Exception as e:
        return templates.TemplateResponse(request, "leads/import.html", {
            "step": "upload", "result": {"error": f"Не удалось прочитать файл: {e}"}})
    if not rows:
        return templates.TemplateResponse(request, "leads/import.html", {
            "step": "upload", "result": {"error": "В файле не найдено строк с данными."}})

    os.makedirs(_TMP_DIR, exist_ok=True)
    token = uuid.uuid4().hex
    with open(os.path.join(_TMP_DIR, token + ".json"), "w", encoding="utf-8") as f:
        json.dump({"filename": file.filename, "headers": headers, "rows": rows},
                  f, ensure_ascii=False)

    cols = _match_columns(headers)
    return templates.TemplateResponse(request, "leads/import.html", {
        "step": "preview", "token": token, "filename": file.filename,
        "headers": headers, "sample": rows[:6], "total": len(rows),
        "guess": cols, "fields": MAPPABLE_FIELDS, "result": None,
    })


@router.post("/import/commit", response_class=HTMLResponse)
@login_required
async def import_commit(request: Request, db: Session = Depends(get_db)):
    """Шаг 2 — применяем маппинг (с правками пользователя) и дедуп, импортируем."""
    form = await request.form()
    token = form.get("token", "")
    dedup = form.get("dedup") == "on"
    path = os.path.join(_TMP_DIR, re.sub(r"\W", "", token) + ".json")
    if not token or not os.path.exists(path):
        return templates.TemplateResponse(request, "leads/import.html", {
            "step": "upload", "result": {"error": "Сессия импорта истекла, загрузите файл заново."}})

    with open(path, encoding="utf-8") as f:
        payload = json.load(f)
    headers, rows, filename = payload["headers"], payload["rows"], payload["filename"]

    # Маппинг с учётом ручных правок: map_<field> = индекс колонки или ""
    cols = {}
    for field in MAPPABLE_FIELDS:
        raw = form.get(f"map_{field}", "")
        if raw != "" and raw is not None:
            try:
                cols[field] = int(raw)
            except ValueError:
                pass

    parsed = [p for cells in rows if (p := _build_lead(cells, cols, headers))]

    # Существующие ключи для дедупликации
    existing_phones, existing_names = set(), set()
    if dedup:
        for ph, nm, addr in db.query(SalesLead.phone, SalesLead.name, SalesLead.address)\
                              .filter(SalesLead.is_active == True).all():
            if ph:
                existing_phones.add(_norm_phone(ph))
            if addr:   # ключ имя+адрес только при заполненном адресе
                existing_names.add((nm or "").lower().strip() + "|" + addr.lower().strip())

    counts: dict[str, int] = {}
    for p in parsed:
        if p["brand"]:
            counts[p["brand"]] = counts.get(p["brand"], 0) + 1

    created, skipped = 0, 0
    seen_phones, seen_names = set(), set()
    networks = set()
    for p in parsed:
        if dedup:
            ph = _norm_phone(p["phone"])
            nk = (p["name"].lower().strip() + "|" + p["address"].lower().strip()
                  if p["address"] else None)
            dup = (ph and (ph in existing_phones or ph in seen_phones)) or \
                  (nk is not None and (nk in existing_names or nk in seen_names))
            if dup:
                skipped += 1
                continue
            if ph:
                seen_phones.add(ph)
            if nk is not None:
                seen_names.add(nk)
        size = counts.get(p["brand"], 1)
        is_net = bool(p["brand"]) and size >= 2
        if is_net:
            networks.add(p["brand"])
        db.add(SalesLead(
            **{k: p[k] for k in (
                "name", "brand", "category", "city", "district", "address", "phone", "email",
                "contact_person", "website", "vk", "instagram", "telegram", "whatsapp", "raw"
            )},
            is_network=is_net, network_size=size,
            source_file=(filename or "import")[:300],
        ))
        created += 1
    db.commit()
    try:
        os.remove(path)
    except OSError:
        pass

    net_points = sum(1 for p in parsed if p["brand"] and counts.get(p["brand"], 1) >= 2)
    result = {
        "created": created, "skipped": skipped,
        "networks": len(networks), "network_points": net_points,
        "singles": created - net_points if not dedup else None,
        "socials": sum(1 for p in parsed if any(p[k] for k in ("vk", "instagram", "telegram", "whatsapp", "website"))),
        "mapped": {MAPPABLE_FIELDS[k]: headers[v] for k, v in cols.items() if v < len(headers)},
        "filename": filename,
    }
    return templates.TemplateResponse(request, "leads/import.html",
                                      {"step": "done", "result": result})


@router.post("/{lead_id}/status", response_class=JSONResponse)
@login_required
async def set_status(request: Request, lead_id: int, status: str = Form(...),
                     comment: str = Form(default=""),
                     db: Session = Depends(get_db)):
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return JSONResponse({"error": "not found"}, status_code=404)
    if status not in LEAD_STATUSES:
        return JSONResponse({"error": "bad status"}, status_code=400)
    lead.call_count = (lead.call_count or 0) + 1
    lead.call_status = status
    lead.last_call_at = datetime.now()
    db.add(LeadCall(lead_id=lead.id, user_id=request.session.get("user_id"),
                    status=status, comment=comment or None))
    db.commit()
    if status == "deal":
        from app.models import CompanySettings
        from app.services.bitrix_client import push_lead_deal_to_bitrix
        push_lead_deal_to_bitrix(lead, db.query(CompanySettings).first(), db=db)
    return JSONResponse({
        "ok": True, "status": status, "label": LEAD_STATUSES[status],
        "color": STATUS_COLORS[status],
        "last_call": lead.last_call_at.strftime("%d.%m.%Y %H:%M"),
    })


@router.get("/{lead_id}/history", response_class=JSONResponse)
@login_required
async def lead_history(request: Request, lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return JSONResponse({"error": "not found"}, status_code=404)
    return JSONResponse({
        "name": lead.name,
        "converted_cp_id": lead.converted_cp_id,
        "calls": [{
            "status": c.status,
            "label": LEAD_STATUSES.get(c.status, c.status or "—"),
            "color": STATUS_COLORS.get(c.status, "secondary"),
            "comment": c.comment or "",
            "user": c.user.full_name if c.user else "—",
            "time": c.created_at.strftime("%d.%m.%Y %H:%M") if c.created_at else "",
        } for c in lead.calls],
    })


@router.post("/{lead_id}/to-counterparty")
@login_required
async def to_counterparty(request: Request, lead_id: int,
                          inn: str = Form(default=""), db: Session = Depends(get_db)):
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return RedirectResponse(url="/leads/", status_code=302)
    if lead.converted_cp_id:
        return RedirectResponse(url=f"/counterparties/{lead.converted_cp_id}", status_code=302)
    cp = Counterparty(
        name=lead.name, short_name=lead.name[:100],
        inn=(inn.strip() or None), phone=lead.phone, email=lead.email,
        actual_address=(f"{lead.city}, {lead.address}" if lead.city else lead.address) or None,
        contact_person=lead.contact_person, type="client",
        notes=f"Создан из прозвона. Рубрика: {lead.category or '—'}."
              + (f" Сайт: {lead.website}" if lead.website else ""),
    )
    db.add(cp)
    db.flush()
    lead.converted_cp_id = cp.id
    if lead.call_status not in ("deal",):
        lead.call_status = "interested"
    db.commit()
    return RedirectResponse(url=f"/counterparties/{cp.id}/edit", status_code=302)


@router.post("/{lead_id}/field", response_class=JSONResponse)
@login_required
async def set_field(request: Request, lead_id: int,
                    field: str = Form(...), value: str = Form(default=""),
                    db: Session = Depends(get_db)):
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return JSONResponse({"error": "not found"}, status_code=404)
    if field == "notes":
        lead.notes = value or None
    elif field == "contact_person":
        lead.contact_person = value or None
    elif field == "phone":
        lead.phone = value or None
    elif field == "district":
        lead.district = value.strip()[:150] or None
    elif field == "assigned_to_id":
        lead.assigned_to_id = int(value) if value else None
    elif field == "callback_at":
        lead.callback_at = date.fromisoformat(value) if value else None
    else:
        return JSONResponse({"error": "bad field"}, status_code=400)
    db.commit()
    return JSONResponse({"ok": True})


@router.post("/{lead_id}/delete")
@role_required("manager")
async def delete_lead(request: Request, lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if lead:
        lead.is_active = False
        db.commit()
    return RedirectResponse(url="/leads/", status_code=302)


@router.post("/clear")
@role_required("manager")
async def clear_source(request: Request, source: str = Form(default=""),
                       db: Session = Depends(get_db)):
    q = db.query(SalesLead).filter(SalesLead.is_active == True)
    if source:
        q = q.filter(SalesLead.source_file == source)
    q.update({SalesLead.is_active: False})
    db.commit()
    return RedirectResponse(url="/leads/", status_code=302)


@router.post("/bulk")
@role_required("manager")
async def bulk_action(request: Request, db: Session = Depends(get_db)):
    """Массовые действия над выбранными точками."""
    form = await request.form()
    ids = [int(x) for x in form.getlist("ids") if x.isdigit()]
    action = form.get("action", "")
    if not ids:
        return RedirectResponse(url="/leads/", status_code=302)
    leads = db.query(SalesLead).filter(SalesLead.id.in_(ids)).all()

    if action == "status":
        status = form.get("value", "")
        if status in LEAD_STATUSES:
            for l in leads:
                l.call_status = status
    elif action == "assign":
        uid = form.get("value", "")
        uid = int(uid) if uid.isdigit() else None
        for l in leads:
            l.assigned_to_id = uid
    elif action == "delete":
        for l in leads:
            l.is_active = False
    elif action == "distribute":
        # равномерно раздать выбранным менеджерам
        mgr_ids = [int(x) for x in form.getlist("managers") if x.isdigit()]
        if mgr_ids:
            for i, l in enumerate(leads):
                l.assigned_to_id = mgr_ids[i % len(mgr_ids)]
    db.commit()
    if action == "status" and form.get("value", "") == "deal":
        from app.models import CompanySettings
        from app.services.bitrix_client import push_lead_deal_to_bitrix
        company = db.query(CompanySettings).first()
        for l in leads:
            push_lead_deal_to_bitrix(l, company, db=db)
    back = request.headers.get("referer", "/leads/")
    return RedirectResponse(url=back, status_code=302)


@router.get("/stats", response_class=HTMLResponse)
@login_required
async def stats_page(request: Request, db: Session = Depends(get_db)):
    base = db.query(SalesLead).filter(SalesLead.is_active == True)
    # Воронка по статусам
    funnel = {}
    for st in LEAD_STATUSES:
        funnel[st] = base.filter(SalesLead.call_status == st).count()
    total = sum(funnel.values())

    # По менеджерам (отдел продаж)
    users = db.query(User).filter(User.is_active == True, User.role == "sales").all()
    by_manager = []
    for u in users:
        uq = base.filter(SalesLead.assigned_to_id == u.id)
        assigned = uq.count()
        if assigned == 0:
            continue
        called = uq.filter(SalesLead.call_status != "new").count()
        interested = uq.filter(SalesLead.call_status.in_(["interested", "thinking"])).count()
        deals = uq.filter(SalesLead.call_status == "deal").count()
        by_manager.append({
            "name": u.full_name, "assigned": assigned, "called": called,
            "interested": interested, "deals": deals,
            "conv": round(deals / called * 100) if called else 0,
        })
    by_manager.sort(key=lambda m: m["deals"], reverse=True)

    return templates.TemplateResponse(request, "leads/stats.html", {
        "funnel": funnel, "total": total, "statuses": LEAD_STATUSES,
        "status_colors": STATUS_COLORS, "by_manager": by_manager,
        "networks": base.filter(SalesLead.is_network == True).count(),
        "singles": base.filter(SalesLead.is_network == False).count(),
        "converted": base.filter(SalesLead.converted_cp_id.isnot(None)).count(),
    })


@router.get("/export.csv")
@login_required
async def export_csv(
    request: Request,
    q: str = "", status: str = "", kind: str = "",
    category: str = "", assigned: str = "", source: str = "",
    db: Session = Depends(get_db),
):
    query = db.query(SalesLead).filter(SalesLead.is_active == True)
    if status:
        query = query.filter(SalesLead.call_status == status)
    if kind == "network":
        query = query.filter(SalesLead.is_network == True)
    elif kind == "single":
        query = query.filter(SalesLead.is_network == False)
    if source:
        query = query.filter(SalesLead.source_file == source)
    if category:
        query = query.filter(SalesLead.category.ilike(f"%{category}%"))
    leads = query.options(joinedload(SalesLead.assigned_to))\
        .order_by(SalesLead.is_network.desc(), SalesLead.brand, SalesLead.name).all()

    buf = io.StringIO()
    buf.write("﻿")  # BOM для Excel
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Название", "Сеть", "Размер сети", "Рубрика", "Город", "Район", "Адрес",
                "Телефон", "Email", "Контакт", "Сайт", "VK", "Instagram",
                "Telegram", "WhatsApp", "Статус", "Менеджер", "Перезвон", "Заметки"])
    for l in leads:
        w.writerow([
            l.name, "Сеть" if l.is_network else "Одиночка", l.network_size or 1,
            l.category or "", l.city or "", l.district or "", l.address or "", l.phone or "",
            l.email or "", l.contact_person or "", l.website or "", l.vk or "",
            l.instagram or "", l.telegram or "", l.whatsapp or "",
            LEAD_STATUSES.get(l.call_status, l.call_status),
            l.assigned_to.full_name if l.assigned_to else "",
            l.callback_at.strftime("%d.%m.%Y") if l.callback_at else "",
            (l.notes or "").replace("\n", " "),
        ])
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=prozvon.csv"},
    )


@router.get("/export.xlsx")
@login_required
async def export_xlsx(
    request: Request,
    status: str = "", kind: str = "", category: str = "", source: str = "",
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    query = db.query(SalesLead).filter(SalesLead.is_active == True)
    if status:
        query = query.filter(SalesLead.call_status == status)
    if kind == "network":
        query = query.filter(SalesLead.is_network == True)
    elif kind == "single":
        query = query.filter(SalesLead.is_network == False)
    if source:
        query = query.filter(SalesLead.source_file == source)
    if category:
        query = query.filter(SalesLead.category.ilike(f"%{category}%"))
    leads = query.options(joinedload(SalesLead.assigned_to))\
        .order_by(SalesLead.is_network.desc(), SalesLead.brand, SalesLead.name).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Прозвон"
    head = ["Название", "Сеть", "Размер", "Рубрика", "Город", "Адрес", "Телефон",
            "Email", "Контакт", "Сайт", "VK", "Instagram", "Telegram", "WhatsApp",
            "Статус", "Менеджер", "Перезвон", "Заметки"]
    ws.append(head)
    fill = PatternFill("solid", fgColor="1E3A5F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    for l in leads:
        ws.append([
            l.name, "Сеть" if l.is_network else "Одиночка", l.network_size or 1,
            l.category or "", l.city or "", l.address or "", l.phone or "",
            l.email or "", l.contact_person or "", l.website or "", l.vk or "",
            l.instagram or "", l.telegram or "", l.whatsapp or "",
            LEAD_STATUSES.get(l.call_status, l.call_status),
            l.assigned_to.full_name if l.assigned_to else "",
            l.callback_at.strftime("%d.%m.%Y") if l.callback_at else "",
            (l.notes or "").replace("\n", " "),
        ])
    widths = [32, 9, 7, 16, 14, 32, 16, 22, 18, 24, 22, 22, 22, 18, 16, 16, 11, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=prozvon.xlsx"},
    )


# ── Карта ────────────────────────────────────────────────────────────────────

# Состояние фонового геокодирования (singleton, один сервер)
_geo_state: dict = {"running": False, "done": 0, "total": 0, "errors": 0, "last_error": None, "blocked": False}
# ID лидов, для которых геокодирование уже пробовали и Nominatim ОКОНЧАТЕЛЬНО ничего не нашёл
# (адрес не существует) — сбрасывается при рестарте сервера, либо вручную через /retry-failed.
# НЕ используется для транзитных ошибок запроса (429/5xx/сеть) — те просто повторяются в
# следующем запуске, иначе один сбой Nominatim навсегда "хоронит" адрес.
_geocode_failed_ids: set[int] = set()


@router.get("/map", response_class=HTMLResponse)
@login_required
async def leads_map(request: Request, db: Session = Depends(get_db)):
    from sqlalchemy import or_
    base = db.query(SalesLead).filter(SalesLead.is_active == True)
    # Считаем только геокодируемые (есть адрес или город) — остальные никогда не получат координаты
    from sqlalchemy import and_
    geocodable = base.filter(
        or_(
            and_(SalesLead.city.isnot(None), SalesLead.city != ""),
            and_(SalesLead.address.isnot(None), SalesLead.address != ""),
        )
    )
    geocoded = geocodable.filter(SalesLead.lat.isnot(None)).count()
    # Ещё не геокодированные, исключая те, которые уже пробовали — Nominatim не нашёл адрес
    remaining_q = geocodable.filter(SalesLead.lat.is_(None))
    if _geocode_failed_ids:
        remaining_q = remaining_q.filter(~SalesLead.id.in_(_geocode_failed_ids))
    remaining = remaining_q.count()
    total = geocoded + remaining
    users = db.query(User).filter(User.is_active == True, User.role == "sales").order_by(User.full_name).all()
    return templates.TemplateResponse(request, "leads/map.html", {
        "total": total, "geocoded": geocoded,
        "users": users,
        "statuses": LEAD_STATUSES, "status_colors": STATUS_COLORS,
        "geo_state": dict(_geo_state),
    })


@router.get("/map/data", response_class=JSONResponse)
@login_required
async def map_data(request: Request, db: Session = Depends(get_db)):
    """Возвращает все геокодированные лиды для отображения на карте."""
    rows = db.query(SalesLead).filter(
        SalesLead.is_active == True,
        SalesLead.lat.isnot(None),
        SalesLead.lng.isnot(None),
    ).options(joinedload(SalesLead.assigned_to)).all()
    return [
        {
            "id": l.id,
            "name": l.name,
            "lat": l.lat,
            "lng": l.lng,
            "status": l.call_status,
            "phone": l.phone or "",
            "category": l.category or "",
            "district": l.district or "",
            "city": l.city or "",
            "address": l.address or "",
            "assigned": l.assigned_to.full_name if l.assigned_to else "",
            "is_network": l.is_network,
            "callback_at": l.callback_at.isoformat() if l.callback_at else None,
        }
        for l in rows
    ]


@router.post("/admin/sync-orders", response_class=JSONResponse)
@role_required("manager")
async def sync_from_orders(request: Request, db: Session = Depends(get_db)):
    """Находит лидов с адресами из заказов и ставит им статус 'deal'."""
    from app.models import Order

    order_addresses = [
        row[0].lower().strip()
        for row in db.query(Order.delivery_address)
            .filter(Order.delivery_address.isnot(None), Order.delivery_address != "")
            .distinct().all()
        if row[0] and row[0].strip()
    ]

    if not order_addresses:
        return JSONResponse({"ok": True, "updated": 0, "msg": "Нет адресов в заказах"})

    leads = db.query(SalesLead).filter(
        SalesLead.is_active == True,
        SalesLead.address.isnot(None),
        SalesLead.address != "",
        SalesLead.call_status != "deal",
    ).all()

    updated = 0
    matched_leads = []
    for lead in leads:
        lead_addr = lead.address.lower().strip()
        if len(lead_addr) < 5:
            continue
        for order_addr in order_addresses:
            if lead_addr in order_addr or order_addr in lead_addr:
                lead.call_status = "deal"
                lead.last_call_at = datetime.now()
                db.add(LeadCall(
                    lead_id=lead.id,
                    status="deal",
                    comment="Авто: адрес совпадает с адресом доставки в заказе",
                ))
                updated += 1
                matched_leads.append(lead)
                break

    db.commit()
    if matched_leads:
        from app.models import CompanySettings
        from app.services.bitrix_client import push_lead_deal_to_bitrix
        company = db.query(CompanySettings).first()
        for lead in matched_leads:
            push_lead_deal_to_bitrix(lead, company, db=db)
    return JSONResponse({"ok": True, "updated": updated})


@router.post("/admin/delete-no-address", response_class=JSONResponse)
@role_required("manager")
async def delete_no_address(request: Request, db: Session = Depends(get_db)):
    """Удаляет (деактивирует) лидов без адреса."""
    from sqlalchemy import or_
    q = db.query(SalesLead).filter(
        SalesLead.is_active == True,
        or_(SalesLead.address.is_(None), SalesLead.address == ""),
    )
    count = q.count()
    q.update({SalesLead.is_active: False}, synchronize_session=False)
    db.commit()
    return JSONResponse({"ok": True, "deleted": count})


@router.get("/geocode/status", response_class=JSONResponse)
@login_required
async def geocode_status(request: Request):
    return dict(_geo_state)


@router.post("/geocode", response_class=JSONResponse)
@login_required
async def geocode_leads(request: Request, db: Session = Depends(get_db)):
    """Запускает фоновое геокодирование точек без координат."""
    if _geo_state["running"]:
        return JSONResponse({"error": "already_running", **_geo_state}, status_code=409)

    pending_q = db.query(SalesLead.id, SalesLead.city, SalesLead.address).filter(
        SalesLead.is_active == True,
        SalesLead.lat.is_(None),
    )
    if _geocode_failed_ids:
        pending_q = pending_q.filter(~SalesLead.id.in_(_geocode_failed_ids))
    pending = pending_q.all()

    if not pending:
        return JSONResponse({"ok": True, "done": 0, "total": 0})

    items = [
        (row.id, f"{row.city or ''} {row.address or ''}".strip())
        for row in pending
        if (row.city or row.address)
    ]

    def _run():
        import logging
        from app.database import SessionLocal
        from app.utils.geocode import GeocodeRequestError, geocode_address_sync
        logger = logging.getLogger("tms.geocode")
        _geo_state.update({
            "running": True, "done": 0, "total": len(items), "errors": 0,
            "last_error": None, "blocked": False,
        })
        s = SessionLocal()
        consecutive_request_errors = 0
        try:
            for i, (lead_id, query) in enumerate(items):
                try:
                    result = geocode_address_sync(query)
                except GeocodeRequestError as e:
                    consecutive_request_errors += 1
                    _geo_state["errors"] += 1
                    _geo_state["last_error"] = str(e)
                    logger.warning("Geocode request failed for lead %s: %s", lead_id, e)
                    # Не баним lead_id навсегда — это сбой запроса, не "адрес не найден".
                    if consecutive_request_errors >= 3:
                        # Похоже, Nominatim блокирует/троттлит нас — прекращаем, чтобы не
                        # прожечь остаток очереди впустую; попробуем снова в следующий запуск.
                        _geo_state["blocked"] = True
                        break
                    time.sleep(e.retry_after or 5)
                    continue
                consecutive_request_errors = 0
                if result:
                    lead = s.get(SalesLead, lead_id)
                    if lead:
                        lead.lat, lead.lng = result
                        s.commit()
                else:
                    _geo_state["errors"] += 1
                    _geocode_failed_ids.add(lead_id)  # адрес окончательно не найден
                _geo_state["done"] = i + 1
                if i < len(items) - 1:
                    time.sleep(1.15)
        except Exception as e:
            logger.exception("Geocode batch crashed")
            _geo_state["last_error"] = str(e)
        finally:
            s.close()
            _geo_state["running"] = False

    threading.Thread(target=_run, daemon=True).start()
    return JSONResponse({"ok": True, "started": True, "total": len(items)})


@router.post("/geocode/retry-failed", response_class=JSONResponse)
@login_required
async def geocode_retry_failed(request: Request):
    """Сбрасывает список адресов, для которых Nominatim ничего не нашёл, чтобы попробовать снова
    (например, после исправления опечатки в адресе или снятия блокировки Nominatim)."""
    count = len(_geocode_failed_ids)
    _geocode_failed_ids.clear()
    return JSONResponse({"ok": True, "cleared": count})
