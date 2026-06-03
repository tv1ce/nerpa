"""Отдел продаж → «Разведка ЛПР».

Отдельный раздел Прозвона: автоматический разбор/обогащение базы лидов и
поиск лица, принимающего решение (ЛПР), по открытым источникам РФ.
"""
from datetime import datetime

import httpx
from fastapi import APIRouter, Request, Depends, Form
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

from app.database import get_db
from app.auth import login_required, role_required
from app.models import SalesLead, LeadCall, CompanySettings
from app.utils import recon as R

router = APIRouter(prefix="/recon", tags=["recon"])
templates = Jinja2Templates(directory="app/templates")


def _settings(db: Session) -> CompanySettings:
    return db.query(CompanySettings).first()


def _active(db: Session):
    return db.query(SalesLead).filter(SalesLead.is_active == True)


@router.get("/", response_class=HTMLResponse)
@login_required
async def index(request: Request, q: str = "", view: str = "", reviewed: str = "",
                source: str = "", sort: str = "", group: str = "",
                db: Session = Depends(get_db)):
    base = _active(db)
    stats = {
        "total": base.count(),
        "with_inn": base.filter(SalesLead.inn.isnot(None), SalesLead.inn != "").count(),
        "with_lpr": base.filter(SalesLead.director.isnot(None), SalesLead.director != "").count(),
        "with_socials": base.filter(
            (SalesLead.vk.isnot(None)) | (SalesLead.instagram.isnot(None)) |
            (SalesLead.telegram.isnot(None)) | (SalesLead.whatsapp.isnot(None)) |
            (SalesLead.website.isnot(None))
        ).count(),
        "reviewed": base.filter(SalesLead.recon_reviewed == True).count(),
        "dead": base.filter(SalesLead.company_status.in_(R.DEAD_STATUSES)).count(),
    }
    settings = _settings(db)
    token = R.get_dadata_token(settings)
    sources = [s[0] for s in _active(db).with_entities(SalesLead.source_file)
               .filter(SalesLead.source_file.isnot(None)).distinct().all()]

    # ── Полный список с фильтрами ──
    rows = base
    if q:
        like = f"%{q}%"
        rows = rows.filter(
            SalesLead.name.ilike(like) | SalesLead.director.ilike(like) |
            SalesLead.inn.ilike(like) | SalesLead.city.ilike(like)
        )
    if view == "lpr":
        rows = rows.filter(SalesLead.director.isnot(None), SalesLead.director != "")
    elif view == "nolpr":
        rows = rows.filter((SalesLead.director.is_(None)) | (SalesLead.director == ""))
    elif view == "dead":
        rows = rows.filter(SalesLead.company_status.in_(R.DEAD_STATUSES))
    if reviewed == "yes":
        rows = rows.filter(SalesLead.recon_reviewed == True)
    elif reviewed == "no":
        rows = rows.filter((SalesLead.recon_reviewed == False) | (SalesLead.recon_reviewed.is_(None)))
    if source:
        rows = rows.filter(SalesLead.source_file == source)
    leads = rows.order_by(SalesLead.recon_reviewed.asc(),
                          SalesLead.enriched_at.desc().nullslast(),
                          SalesLead.name).all()

    # Скоринг «теплота»
    scored = [(l, R.lead_score(l)) for l in leads]
    if sort == "hot":
        scored.sort(key=lambda x: x[1], reverse=True)
    leads_scored = [{"lead": l, "score": s, "label": R.score_label(s)} for l, s in scored]

    # ── Группировка по компаниям (одни реквизиты = одна карточка) ──
    companies = None
    if group == "company":
        from collections import OrderedDict
        buckets: "OrderedDict[str, list]" = OrderedDict()
        for l, s in scored:
            buckets.setdefault(R.network_key(l) or f"id:{l.id}", []).append((l, s))
        companies = []
        for key, items in buckets.items():
            head = items[0][0]
            companies.append({
                "key": key, "head": head, "count": len(items),
                "points": [{"lead": l, "score": s} for l, s in items],
                "score": max(s for _, s in items),
                "cities": sorted({l.city for l, _ in items if l.city}),
                "director": head.director, "inn": head.inn,
                "company_status": head.company_status,
                "called": sum(1 for l, _ in items if (l.call_status or "new") != "new"),
            })
        companies.sort(key=lambda c: (-c["count"], c["head"].name))

    duplicates = sum(len(d) for _, d in R.find_duplicate_groups(leads)) if leads else 0

    return templates.TemplateResponse(request, "recon/index.html", {
        "stats": stats, "has_token": bool(token), "sources": sources,
        "leads": leads_scored, "companies": companies, "duplicates": duplicates,
        "status_badge": R.STATUS_BADGE,
        "statuses": {  # для отображения статуса прозвона
            "new": "Новый", "callback": "Перезвонить", "no_answer": "Не дозвонился",
            "interested": "Заинтересован", "thinking": "Думает", "refused": "Отказ",
            "deal": "Продажа", "invalid": "Невалид",
        },
        "q": q, "view": view, "reviewed": reviewed, "source": source,
        "sort": sort, "group": group,
    })


@router.post("/config")
@login_required
async def save_config(request: Request, dadata_token: str = Form(default=""),
                      dadata_secret: str = Form(default=""), db: Session = Depends(get_db)):
    settings = _settings(db)
    if settings:
        settings.dadata_token = dadata_token.strip() or None
        settings.dadata_secret = dadata_secret.strip() or None
        db.commit()
    return RedirectResponse(url="/recon/", status_code=302)


@router.post("/run-local", response_class=JSONResponse)
@role_required("manager")
async def run_local(request: Request, source: str = Form(default=""),
                    db: Session = Depends(get_db)):
    """Локальное обогащение всей базы: соцсети/почта/телефон/ИНН + пересборка сетей."""
    q = _active(db)
    if source:
        q = q.filter(SalesLead.source_file == source)
    leads = q.all()

    field_hits: dict[str, int] = {}
    touched = 0
    for l in leads:
        changed = R.enrich_lead_local(l)
        if changed:
            touched += 1
            for f in changed:
                field_hits[f] = field_hits.get(f, 0) + 1
    nets = R.reclassify_networks(leads)
    db.commit()

    labels = {
        "vk": "VK", "instagram": "Instagram", "telegram": "Telegram",
        "whatsapp": "WhatsApp", "website": "Сайт", "email": "E-mail",
        "phone": "Телефон", "phone_fmt": "Формат телефона", "inn": "ИНН",
        "brand": "Бренд (для сетей)",
    }
    return JSONResponse({
        "ok": True, "processed": len(leads), "touched": touched,
        "networks_changed": nets,
        "fields": {labels.get(k, k): v for k, v in sorted(field_hits.items(), key=lambda x: -x[1])},
    })


@router.post("/run-dadata", response_class=JSONResponse)
@role_required("manager")
async def run_dadata(request: Request, source: str = Form(default=""),
                     limit: int = Form(default=50), only_missing: str = Form(default="on"),
                     drop_dead: str = Form(default="on"),
                     db: Session = Depends(get_db)):
    """Пакетное обогащение через DaData (ЛПР + реквизиты). Ограничено лимитом.
    drop_dead=on — ликвидирующиеся/банкроты автоматически убираются из базы."""
    token = R.get_dadata_token(_settings(db))
    if not token:
        return JSONResponse({"error": "Не задан ключ DaData. Укажите его в настройках раздела."}, status_code=400)

    q = _active(db)
    if source:
        q = q.filter(SalesLead.source_file == source)
    if only_missing == "on":
        q = q.filter((SalesLead.director.is_(None)) | (SalesLead.director == ""))
    leads = q.order_by(SalesLead.is_network.desc(), SalesLead.brand).limit(max(1, min(limit, 300))).all()

    found, errors, removed, cached = 0, 0, 0, 0
    # Кэш запросов: один и тот же запрос (ИНН или название+город) спрашиваем у DaData
    # только один раз — для сети/дублей переиспользуем результат (экономим лимит и склеиваем).
    cache: dict[str, dict | None] = {}
    for l in leads:
        try:
            key = l.inn if (l.inn and R._valid_inn(l.inn)) else \
                " ".join(p for p in (l.name, l.city) if p).lower()
            if key in cache:
                res = cache[key]
                if res:
                    R.apply_party(l, res)   # дубль/точка сети — без нового запроса
                    cached += 1
            else:
                res = await R.enrich_lead_dadata(token, l)
                cache[key] = res
            if res and res.get("director"):
                found += 1
            # Авто-удаление ликвидирующихся / банкротов
            if drop_dead == "on" and l.company_status in R.DEAD_STATUSES:
                l.is_active = False
                if l.call_status == "new":
                    l.call_status = "invalid"
                removed += 1
        except httpx.HTTPStatusError as e:
            errors += 1
            if e.response.status_code in (401, 403):
                db.commit()
                return JSONResponse({"error": "DaData отклонила ключ (401/403). Проверьте токен."}, status_code=400)
        except Exception:
            errors += 1

    # Пересобираем сети по реквизитам (одинаковый ИНН → одна сеть)
    db.flush()   # autoflush=False — учитываем деактивацию ликвидированных
    R.reclassify_networks(_active(db).all())
    db.commit()
    return JSONResponse({
        "ok": True, "processed": len(leads), "found_lpr": found,
        "removed": removed, "from_cache": cached, "errors": errors,
    })


@router.post("/cleanup-dead", response_class=JSONResponse)
@role_required("manager")
async def cleanup_dead(request: Request, db: Session = Depends(get_db)):
    """Разовая чистка: убрать из активной базы все уже помеченные как
    ликвидирующиеся / ликвидированные / банкроты (по ранее полученному статусу)."""
    leads = _active(db).filter(SalesLead.company_status.in_(R.DEAD_STATUSES)).all()
    for l in leads:
        l.is_active = False
        if l.call_status == "new":
            l.call_status = "invalid"
    db.commit()
    return JSONResponse({"ok": True, "removed": len(leads)})


@router.get("/dedup-preview", response_class=JSONResponse)
@login_required
async def dedup_preview(request: Request, db: Session = Depends(get_db)):
    """Сколько истинных дублей нашлось (для кнопки, без изменений)."""
    groups = R.find_duplicate_groups(_active(db).all())
    dupes = sum(len(d) for _, d in groups)
    return JSONResponse({"ok": True, "groups": len(groups), "duplicates": dupes})


@router.post("/dedup", response_class=JSONResponse)
@role_required("manager")
async def dedup(request: Request, db: Session = Depends(get_db)):
    """Склеивает истинные дубли (одинаковый ИНН+адрес или телефон):
    оставляет самую полную запись, переносит на неё историю звонков и лучший
    статус прозвона, остальные дубли деактивирует. Сети при этом не трогаются —
    объединяются только точки-близнецы."""
    groups = R.find_duplicate_groups(_active(db).all())
    merged = 0
    for survivor, dupes in groups:
        for d in dupes:
            # переносим историю звонков
            db.query(LeadCall).filter(LeadCall.lead_id == d.id)\
                .update({LeadCall.lead_id: survivor.id})
            # сохраняем лучший статус прозвона
            if R._STATUS_RANK.get(d.call_status or "new", 0) > \
               R._STATUS_RANK.get(survivor.call_status or "new", 0):
                survivor.call_status = d.call_status
            survivor.call_count = (survivor.call_count or 0) + (d.call_count or 0)
            # добиваем пустые поля выжившего из дубля
            for f in ("phone", "email", "director", "director_post", "inn", "ogrn",
                      "address", "vk", "instagram", "telegram", "whatsapp", "website",
                      "contact_person", "okved", "company_status", "company_name_full"):
                if not getattr(survivor, f) and getattr(d, f):
                    setattr(survivor, f, getattr(d, f))
            d.is_active = False
            d.call_status = "invalid"
            merged += 1
    db.flush()   # autoflush=False — фиксируем деактивацию до пересборки сетей
    R.reclassify_networks(_active(db).all())
    db.commit()
    return JSONResponse({"ok": True, "merged": merged, "groups": len(groups)})


@router.get("/export-lpr.xlsx")
@login_required
async def export_lpr(request: Request, db: Session = Depends(get_db)):
    """Выгрузка обогащённой базы с ЛПР и реквизитами в Excel."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    leads = _active(db).filter(SalesLead.director.isnot(None), SalesLead.director != "")\
        .order_by(SalesLead.company_status, SalesLead.name).all()
    wb = Workbook(); ws = wb.active; ws.title = "ЛПР"
    head = ["Точка", "Город", "ЛПР", "Должность", "Телефон", "Email", "ИНН",
            "ОГРН", "Статус ЕГРЮЛ", "ОКВЭД", "Полное наименование", "Температура"]
    ws.append(head)
    fill = PatternFill("solid", fgColor="198754")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = fill
    for l in leads:
        sc = R.lead_score(l)
        ws.append([
            l.name, l.city or "", l.director or "", l.director_post or "",
            R.pretty_phone(l.phone) if l.phone else "", l.email or "", l.inn or "",
            l.ogrn or "", l.company_status or "", l.okved or "",
            l.company_name_full or "", R.score_label(sc)[0],
        ])
    for i, w in enumerate([30, 14, 26, 22, 18, 22, 14, 16, 14, 28, 36, 14], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(
        out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=lpr.xlsx"})


@router.get("/{lead_id:int}", response_class=HTMLResponse)
@login_required
async def lead_card(request: Request, lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return RedirectResponse(url="/recon/", status_code=302)
    # Открытие карточки = «просмотрено»
    if not lead.recon_reviewed:
        lead.recon_reviewed = True
        lead.recon_reviewed_at = datetime.now()
        db.commit()
    links = R.group_links(R.build_source_links(lead))
    token = R.get_dadata_token(_settings(db))
    phones = R.extract_phones(lead)
    return templates.TemplateResponse(request, "recon/lead.html", {
        "lead": lead, "links": links, "has_token": bool(token), "phones": phones,
        "status_badge": R.STATUS_BADGE, "pretty_phone": R.pretty_phone,
    })


@router.post("/{lead_id:int}/reviewed", response_class=JSONResponse)
@login_required
async def toggle_reviewed(request: Request, lead_id: int, db: Session = Depends(get_db)):
    """Ручное переключение отметки «просмотрено»."""
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return JSONResponse({"error": "not found"}, status_code=404)
    lead.recon_reviewed = not lead.recon_reviewed
    lead.recon_reviewed_at = datetime.now() if lead.recon_reviewed else None
    db.commit()
    return JSONResponse({"ok": True, "reviewed": lead.recon_reviewed})


@router.post("/{lead_id:int}/dadata", response_class=JSONResponse)
@login_required
async def lead_dadata(request: Request, lead_id: int, query: str = Form(default=""),
                      db: Session = Depends(get_db)):
    """Поиск в DaData по одной точке. query пустой → по ИНН/названию лида."""
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return JSONResponse({"error": "not found"}, status_code=404)
    token = R.get_dadata_token(_settings(db))
    if not token:
        return JSONResponse({"error": "Не задан ключ DaData."}, status_code=400)

    try:
        if query.strip():
            sugs = await R.dadata_suggest(token, query.strip(), count=8)
            return JSONResponse({"ok": True, "candidates": [R.parse_party(s) for s in sugs]})
        res = await R.enrich_lead_dadata(token, lead)
        db.commit()
        if not res:
            return JSONResponse({"ok": True, "candidates": []})
        return JSONResponse({"ok": True, "applied": True, "data": res})
    except httpx.HTTPStatusError as e:
        msg = "DaData отклонила ключ." if e.response.status_code in (401, 403) else f"Ошибка DaData ({e.response.status_code})."
        return JSONResponse({"error": msg}, status_code=400)
    except Exception as e:
        return JSONResponse({"error": f"Сбой запроса: {e}"}, status_code=500)


@router.post("/{lead_id:int}/apply", response_class=JSONResponse)
@login_required
async def lead_apply(request: Request, lead_id: int, db: Session = Depends(get_db)):
    """Применяет выбранного кандидата DaData (поля приходят формой)."""
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return JSONResponse({"error": "not found"}, status_code=404)
    form = await request.form()
    for field in ("inn", "kpp", "ogrn", "company_name_full", "director",
                  "director_post", "company_status", "okved", "registration_date"):
        val = (form.get(field) or "").strip()
        if val:
            setattr(lead, field, val)
    if lead.director and not lead.contact_person:
        lead.contact_person = lead.director[:150]
    lead.enriched_at = datetime.now()
    lead.enrich_source = "dadata"
    db.commit()
    return JSONResponse({"ok": True})


@router.post("/{lead_id:int}/lpr", response_class=JSONResponse)
@login_required
async def lead_lpr(request: Request, lead_id: int, db: Session = Depends(get_db)):
    """Ручное сохранение ЛПР/реквизитов из карточки."""
    lead = db.query(SalesLead).filter(SalesLead.id == lead_id).first()
    if not lead:
        return JSONResponse({"error": "not found"}, status_code=404)
    form = await request.form()
    mapping = {
        "director": 200, "director_post": 200, "inn": 12, "ogrn": 15,
        "company_name_full": 500, "okved": 300, "contact_person": 150,
    }
    for field, ln in mapping.items():
        if field in form:
            val = (form.get(field) or "").strip()
            setattr(lead, field, val[:ln] or None)
    lead.enrich_source = "manual"
    lead.enriched_at = datetime.now()
    db.commit()
    return JSONResponse({"ok": True})
