"""Заполнение .docx шаблонов данными и конвертация в PDF."""
import os
import shutil
from datetime import date
from docx import Document

# Месяцы в родительном падеже
_MONTHS_RU = {
    1: "января", 2: "февраля", 3: "марта", 4: "апреля",
    5: "мая", 6: "июня", 7: "июля", 8: "августа",
    9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
}

# Числа в родительном падеже (для отсрочки дней)
_DAYS_GENITIVE = {
    1: "одного", 2: "двух", 3: "трёх", 4: "четырёх", 5: "пяти",
    6: "шести", 7: "семи", 8: "восьми", 9: "девяти", 10: "десяти",
    11: "одиннадцати", 12: "двенадцати", 13: "тринадцати", 14: "четырнадцати",
    15: "пятнадцати", 16: "шестнадцати", 17: "семнадцати", 18: "восемнадцати",
    19: "девятнадцати", 20: "двадцати", 21: "двадцати одного", 22: "двадцати двух",
    23: "двадцати трёх", 24: "двадцати четырёх", 25: "двадцати пяти",
    26: "двадцати шести", 27: "двадцати семи", 28: "двадцати восьми",
    29: "двадцати девяти", 30: "тридцати", 31: "тридцати одного",
    45: "сорока пяти", 60: "шестидесяти", 90: "девяноста",
}


def _get_signatory(cp) -> str:
    """Возвращает подписанта контрагента. Для ИП без явного значения генерирует из ФИО."""
    if getattr(cp, "signatory", None):
        return cp.signatory
    if getattr(cp, "entity_type", None) == "ip":
        source = cp.contact_person or cp.name or ""
        if not source:
            return "___________________"
        name = source.strip()
        for prefix in ("Индивидуальный предприниматель ", "индивидуальный предприниматель ", "ИП ", "ип "):
            if name.startswith(prefix):
                name = name[len(prefix):]
                break
        parts = name.split()
        if len(parts) >= 3:
            return f"{parts[0]} {parts[1][0].upper()}.{parts[2][0].upper()}."
        if len(parts) == 2:
            return f"{parts[0]} {parts[1][0].upper()}."
        return parts[0] if parts else "___________________"
    return cp.contact_person or "___________________"


def _fmt_date(d) -> str:
    if not d:
        return ""
    if isinstance(d, date):
        return d.strftime("%d.%m.%Y")
    return str(d)


def _fmt_money(v) -> str:
    if v is None:
        return "0,00"
    return f"{v:,.2f}".replace(",", " ").replace(".", ",")


def _days_words(n) -> str:
    """Число дней прописью в родительном падеже."""
    if n is None:
        return "__________"
    if n in _DAYS_GENITIVE:
        return _DAYS_GENITIVE[n]
    return "__________"


def fill_contract_template(template_path: str, output_path: str, contract, company) -> str:
    """Заполняет .docx шаблон и сохраняет результат. Возвращает путь к файлу."""
    doc = Document(template_path)
    cp = contract.counterparty

    d = contract.date
    placeholders = {
        "{{contract_number}}": contract.number or "",
        "{{contract_date}}": _fmt_date(contract.date),
        "{{contract_date_day}}": str(d.day) if d else "___",
        "{{contract_date_month}}": _MONTHS_RU.get(d.month, "") if d else "___________",
        "{{contract_date_year}}": str(d.year) if d else "____",
        "{{contract_subject}}": contract.subject or "",
        "{{contract_amount}}": _fmt_money(contract.amount),
        "{{start_date}}": _fmt_date(contract.start_date),
        "{{end_date}}": _fmt_date(contract.end_date),
        # Отсрочка платежа
        "{{payment_days_num}}": str(contract.payment_days) if contract.payment_days else "___",
        "{{payment_days_words}}": _days_words(contract.payment_days),
        # Наша компания
        "{{company_name}}": company.name or "",
        "{{company_short_name}}": company.short_name or company.name or "",
        "{{company_inn}}": company.inn or "",
        "{{company_kpp}}": company.kpp or "",
        "{{company_ogrn}}": company.ogrn or "",
        "{{company_okpo}}": getattr(company, "okpo", "") or "",
        "{{company_address}}": company.legal_address or "",
        "{{company_phone}}": company.phone or "",
        "{{company_email}}": company.email or "",
        "{{company_director}}": company.director or "",
        "{{company_director_basis}}": company.director_basis or "Устава",
        "{{company_bank_name}}": company.bank_name or "",
        "{{company_bank_account}}": company.bank_account or "",
        "{{company_bik}}": company.bank_bik or "",
        "{{company_corr_account}}": company.bank_corr_account or "",
        # Контрагент
        "{{client_name}}": cp.name or "",
        "{{client_short_name}}": cp.short_name or cp.name or "",
        "{{client_inn}}": cp.inn or "",
        "{{client_kpp}}": cp.kpp or "",
        "{{client_ogrn}}": cp.ogrn or "",
        "{{client_address}}": cp.legal_address or "",
        "{{client_phone}}": cp.phone or "",
        "{{client_email}}": cp.email or "",
        "{{client_contact}}": cp.contact_person or "",
        "{{client_rep}}": cp.contact_person or "___________________",
        "{{client_signatory}}": _get_signatory(cp),
        "{{client_bank_name}}": cp.bank_name or "",
        "{{client_bank_account}}": cp.bank_account or "",
        "{{client_bik}}": cp.bank_bik or "",
        "{{client_corr_account}}": cp.bank_corr_account or "",
    }

    for para in doc.paragraphs:
        _replace_in_para(para, placeholders)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_para(para, placeholders)

    doc.save(output_path)
    return output_path


def _replace_in_para(para, replacements: dict):
    """Заменяет плейсхолдеры в параграфе.

    Word часто разбивает текст на несколько runs ({{, placeholder, }}),
    поэтому сначала склеиваем всё в строку, делаем замены, затем
    кладём результат в первый run и очищаем остальные.
    Форматирование (шрифт, размер) первого run сохраняется.
    """
    if not para.runs:
        return

    full_text = "".join(run.text for run in para.runs)

    # Проверяем нужна ли вообще хоть одна замена
    if not any(key in full_text for key in replacements):
        return

    new_text = full_text
    for key, value in replacements.items():
        new_text = new_text.replace(key, value)

    # Кладём результат в первый run, остальные обнуляем
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ""


def convert_docx_to_pdf(docx_path: str, output_dir: str) -> str | None:
    """Конвертирует .docx в PDF используя LibreOffice или MS Word (Windows)."""
    try:
        import subprocess
        # Попытка через LibreOffice
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", output_dir, docx_path],
            capture_output=True, timeout=30
        )
        if result.returncode == 0:
            base = os.path.splitext(os.path.basename(docx_path))[0]
            return os.path.join(output_dir, base + ".pdf")
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    try:
        import subprocess
        # Попытка через MS Word на Windows
        script = (
            f"$word = New-Object -ComObject Word.Application;"
            f"$doc = $word.Documents.Open('{docx_path.replace(chr(92), chr(92)*2)}');"
            f"$doc.SaveAs('{os.path.join(output_dir, os.path.splitext(os.path.basename(docx_path))[0] + '.pdf').replace(chr(92), chr(92)*2)}', 17);"
            f"$doc.Close(); $word.Quit()"
        )
        result = subprocess.run(["powershell", "-Command", script], capture_output=True, timeout=60)
        if result.returncode == 0:
            base = os.path.splitext(os.path.basename(docx_path))[0]
            return os.path.join(output_dir, base + ".pdf")
    except Exception:
        pass

    return None
